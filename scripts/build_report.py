#!/usr/bin/env python3
"""Assemble the written report from the artefacts of YOUR run (metrics, figures, screenshots).

    python scripts/build_report.py --author "Your Name" --student-id 2024XXXXX \
        --repo-url https://github.com/<you>/heart-disease-mlops --pdf

Outputs reports/Heart_Disease_MLOps_Report.{docx,md[,pdf]}. One content model feeds both
formats. Screenshots that do not exist yet appear as visible "pending" notes - capture them
(see screenshots/README.md) and re-run. Numbers are read from models/metadata.json, so the
report always matches the model that was actually trained.
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

FIG = ROOT / "reports" / "figures"
MFIG = FIG / "model"
SHOTS = ROOT / "screenshots"
MISSING: list[str] = []

SCREENSHOTS = {
    "mlflow_experiments.png": "MLflow experiment view: parent run with three nested candidate runs",
    "mlflow_compare.png": "MLflow run comparison (parallel coordinates / metric table)",
    "mlflow_run_detail.png": "MLflow single-run page showing parameters, metrics and artifacts",
    "ci_workflow_overview.png": "GitHub Actions: all five jobs green",
    "ci_workflow_artifacts.png": "GitHub Actions: uploaded artifacts (model, mlruns, reports)",
    "ci_failure_demo.png": "GitHub Actions: deliberately failing run (pipeline stops on error)",
    "docker_build.png": "Terminal: docker build output",
    "docker_run_predict.png": "Terminal: container running + successful /predict call",
    "swagger_ui.png": "FastAPI Swagger UI (/docs) with a /predict response",
    "k8s_pods_services.png": "kubectl get pods,svc,hpa,ingress -n heart-risk",
    "k8s_predict_via_lb.png": "curl to the LoadBalancer endpoint returning a prediction",
    "k8s_ingress.png": "Ingress / port-forward access (optional)",
    "prometheus_targets.png": "Prometheus Targets page: API pods UP",
    "grafana_dashboard.png": "Grafana dashboard populated by generate_traffic.py",
    "api_logs.png": "Structured JSON request logs (docker logs / kubectl logs)",
}


# ------------------------------------------------------------------ content model helpers
def H1(t): return ("h1", t)
def H2(t): return ("h2", t)
def P(t): return ("p", t)
def BUL(items): return ("bul", items)
def TAB(head, rows, widths=None): return ("tab", head, rows, widths)
def CODE(t): return ("code", t)
def BR(): return ("br",)


def IMG(path: Path, caption: str, width: float = 5.6):
    if path.exists():
        return ("img", path, caption, width)
    MISSING.append(str(path.relative_to(ROOT)))
    return ("note", f"[Figure pending: {path.name} - {caption}]")


def SHOT(name: str, width: float = 5.2):
    return IMG(SHOTS / name, SHOTS_CAPTION[name], width)


SHOTS_CAPTION = SCREENSHOTS


def load_json(path: Path, default=None):
    return json.loads(path.read_text()) if path.exists() else (default or {})


def fmt(x, nd=3):
    return f"{x:.{nd}f}" if isinstance(x, (int, float)) else str(x)


# ------------------------------------------------------------------ content
def build_content(a) -> list:
    meta = load_json(ROOT / "models" / "metadata.json")
    dq = load_json(ROOT / "data" / "processed" / "data_quality.json")
    if not meta:
        print("WARNING: models/metadata.json not found - run `make train` first", file=sys.stderr)
    all_models = meta.get("all_models", {})
    best = meta.get("model_name", "n/a")
    m = meta.get("metrics", {})
    split = meta.get("split", {})
    libs = meta.get("library_versions", {})
    n_tests = _count_tests()
    top_corr = _top_correlations()
    example = _api_example()

    c: list = []
    # ---- 1
    c += [H1("1. Introduction and problem statement"),
          P("Cardiovascular disease is a leading cause of death worldwide, and early risk "
            "screening from routine clinical measurements can prioritise patients for further "
            "diagnostics. This project builds a binary classifier that predicts the presence of "
            "heart disease from 13 clinical attributes in the UCI Heart Disease (Cleveland) "
            "dataset, and - more importantly for this assignment - wraps the model in a "
            "reproducible, tested, containerised, deployed and monitored service."),
          P("The emphasis is on the MLOps lifecycle: every step from raw data to a monitored "
            "Kubernetes deployment is scripted, version-controlled and re-runnable from a clean "
            "checkout using the pinned requirements files."),
          TAB(["Item", "Value"], [
              ["Dataset", "UCI Heart Disease - Cleveland (processed.cleveland.data)"],
              ["Rows / columns", f"{dq.get('rows', 303)} patients, 13 features + target"],
              ["Task", "Binary classification (1 = heart disease present)"],
              ["Selected model", f"{best} (version {meta.get('model_version', 'n/a')})"],
              ["Stack", "scikit-learn, MLflow, FastAPI, Docker, Kubernetes, GitHub Actions, "
                        "Prometheus, Grafana"],
              ["Repository", a.repo_url]], [1.6, 4.8]),
          P("Disclaimer: this is an educational prototype trained on 303 patients from one "
            "clinic. It is not a medical device and must not be used for clinical decisions.")]

    # ---- 2 architecture
    c += [H1("2. Solution architecture"),
          IMG(ROOT / "docs" / "architecture.png", "End-to-end architecture", 6.4),
          P("Three lanes: (1) data and model development, (2) CI/CD in GitHub Actions that gates "
            "every change through lint, tests, training and a container smoke test, and "
            "(3) the Kubernetes runtime with Prometheus/Grafana observability. The trained "
            "pipeline (preprocessing + model) is a single joblib artefact, so the API applies "
            "exactly the transformations used in training.")]

    # ---- 3 setup
    c += [H1("3. Setup and installation"),
          P("Prerequisites: macOS or Linux, Python 3.11, Git, Docker Desktop, kubectl and "
            "(optionally) Helm and Minikube. On macOS everything can be installed with "
            "scripts/bootstrap_mac.sh."),
          CODE("git clone " + a.repo_url + "\ncd heart-disease-mlops\n"
               "make setup        # venv + pinned dependencies (requirements-dev.txt)\n"
               "make data         # download UCI data and build data/processed/heart_clean.csv\n"
               "make eda          # figures -> reports/figures\n"
               "make train        # tune, track in MLflow, save models/model.joblib\n"
               "make test         # pytest + coverage\n"
               "make docker-build docker-run docker-test\n"
               "make k8s-deploy   # Docker Desktop / Minikube + Prometheus + Grafana"),
          P("The same sequence is executed automatically by scripts/run_pipeline.sh and by the "
            "GitHub Actions workflow, which start from an empty environment - this is the "
            "'clean setup' proof required by the production-readiness criteria.")]

    # ---- 4 data & EDA
    c += [H1("4. Data acquisition and exploratory data analysis"),
          P("scripts/download_data.py fetches the dataset from the UCI repository (zip first, "
            "legacy URL second, ucimlrepo third), validates its shape (303 rows x 14 columns) "
            "and stores it unchanged in data/raw/. heartrisk.data then produces the cleaned CSV "
            "and a data_quality.json report."),
          H2("4.1 Cleaning decisions"),
          BUL(["'?' placeholders became NaN; all columns were coerced to numeric; duplicates "
               "were removed.",
               "The 5-level diagnosis 'num' was binarised: 0 = no disease, 1-4 = disease.",
               f"Missing values: {dq.get('missing_by_column', {})}. They are not imputed in the "
               "CSV; the median/mode imputers live inside the model pipeline and are fitted on "
               "training folds only, preventing leakage.",
               f"Class balance: {dq.get('class_counts', {})} "
               f"(positive rate {fmt(dq.get('positive_rate', 0), 3)})."]),
          H2("4.2 Findings"),
          IMG(FIG / "01_class_balance.png", "Class balance", 3.4),
          IMG(FIG / "03_numeric_histograms.png", "Numeric feature distributions by outcome", 6.2),
          IMG(FIG / "05_categorical_vs_target.png", "Disease rate by categorical feature", 6.2),
          IMG(FIG / "06_correlation_heatmap.png", "Correlation heat-map", 4.8),
          P("Strongest linear associations with the target in this run: " + top_corr + ". "
            "Threshold-style clinical effects and relatively weak individual correlations "
            "motivated adding engineered features and comparing a linear model with tree "
            "ensembles.")]

    # ---- 5 modelling
    rows_cv, rows_test = [], []
    for name, r in all_models.items():
        rows_cv.append([name] + [f"{fmt(r[f'cv_{k}_mean'])} ± {fmt(r[f'cv_{k}_std'], 3)}"
                                 for k in ("accuracy", "precision", "recall", "roc_auc")])
        rows_test.append([name] + [fmt(r[f"test_{k}"]) for k in
                                   ("accuracy", "precision", "recall", "f1", "roc_auc")])
    c += [H1("5. Feature engineering and model development"),
          H2("5.1 Feature pipeline"),
          TAB(["Group", "Columns", "Treatment"], [
              ["Numeric", "age, trestbps, chol, thalach, oldpeak, ca, hr_pct_max",
               "median impute -> StandardScaler"],
              ["Binary", "sex, fbs, exang, high_bp, high_chol, st_depressed", "mode impute"],
              ["Categorical", "cp, restecg, slope, thal", "mode impute -> one-hot "
                                                          "(unknown categories ignored)"],
              ["Engineered", "hr_pct_max = thalach / (220 - age); high_bp >= 140; "
                             "high_chol >= 240; st_depressed = oldpeak > 0",
               "ClinicalFeatureAdder (NaN-safe)"]], [1.0, 3.2, 2.2]),
          P("All steps are scikit-learn transformers inside one Pipeline, so cross-validation "
            "re-fits preprocessing within every fold and the API reuses the identical object."),
          H2("5.2 Models and tuning"),
          P("Three classifiers were compared: Logistic Regression (interpretable baseline), "
            "Random Forest and Gradient Boosting (non-linear ensembles). Each was tuned with an "
            "exhaustive grid search (stratified 5-fold, ROC-AUC objective) on the training "
            f"split ({split.get('train_rows', '?')} rows); "
            f"{split.get('test_rows', '?')} rows were held out untouched."),
          BUL(["Selection protocol: the winning configuration is re-scored on a different CV "
               "shuffle (GridSearchCV.best_score_ is optimistically biased) and the final model "
               "is chosen by cross-validated ROC-AUC - never by test performance.",
               "Recall is monitored closely because a missed patient (false negative) is more "
               "costly than a false alarm in a screening setting."]),
          H2("5.3 Results"),
          P("Cross-validated metrics on the training split (mean ± std):"),
          TAB(["Model", "Accuracy", "Precision", "Recall", "ROC-AUC"], rows_cv),
          P("Hold-out test-set metrics:"),
          TAB(["Model", "Accuracy", "Precision", "Recall", "F1", "ROC-AUC"], rows_test),
          P(f"Selected model: {best} with parameters {meta.get('best_params', {})}. "
            f"Cross-validated ROC-AUC {fmt(m.get('cv_roc_auc_mean', 0))}; "
            f"hold-out ROC-AUC {fmt(m.get('test_roc_auc', 0))}, recall "
            f"{fmt(m.get('test_recall', 0))}, precision {fmt(m.get('test_precision', 0))}. "
            "With only ~60 test patients, these hold-out figures carry wide uncertainty "
            "(a single patient changes recall by roughly 3 points), which is why selection "
            "relies on cross-validation."),
          IMG(MFIG / "cv_model_comparison.png", "Cross-validation comparison", 5.0),
          IMG(MFIG / "roc_overlay_test.png", "ROC curves on the hold-out set", 4.2)]
    if best != "n/a":
        c += [IMG(MFIG / f"confusion_{best}.png", f"Confusion matrix - {best}", 3.2),
              IMG(MFIG / f"importance_{best}.png", f"Feature importance - {best}", 4.4)]

    # ---- 6 tracking
    c += [H1("6. Experiment tracking (MLflow)"),
          P("Every training invocation creates one parent run ('training-session') with one "
            "nested child run per candidate model. Tracking data live in ./mlruns "
            "(browse with `make mlflow-ui`)."),
          TAB(["Logged item", "Where", "Examples"], [
              ["Parameters", "child run", "model_name, best hyper-parameters, grid_size"],
              ["Parameters", "parent run", "seed, split sizes, dataset hash, sklearn version, "
                                           "git commit"],
              ["Metrics", "child run", "cv_{accuracy,precision,recall,f1,roc_auc}_{mean,std}, "
                                       "test_* metrics, grid_best_cv_roc_auc"],
              ["Plots", "child run /plots", "confusion matrix, feature importance"],
              ["Tuning table", "child run /tuning", "full GridSearchCV results (CSV)"],
              ["Model", "child run /model", "sklearn Pipeline + signature + input example"],
              ["Final package", "parent /final_model", "model.joblib, metadata.json, comparison "
                                                       "CSV, ROC overlay, CV bar chart"]],
              [1.2, 1.5, 3.7]),
          SHOT("mlflow_experiments.png"), SHOT("mlflow_compare.png"),
          SHOT("mlflow_run_detail.png")]

    # ---- 7 packaging
    c += [H1("7. Model packaging and reproducibility"),
          BUL(["Deployable artefact: models/model.joblib (full Pipeline) + models/metadata.json "
               "(name, version = semantic version + SHA-256 prefix of the artefact, metrics, "
               "feature list, training-data reference ranges, library versions, git commit).",
               "MLflow model flavour (with signature) is stored for every candidate as a "
               "second reusable format.",
               "requirements-serve.txt (runtime only, used in Docker), requirements.txt "
               "(training + tracking) and requirements-dev.txt (tests, lint, notebooks) all pin "
               "exact versions.",
               "Fixed random seed (default 42, overridable); a unit test asserts that two "
               "training runs produce identical metrics and parameters.",
               "Training environment recorded in metadata: "
               + ", ".join(f"{k} {v}" for k, v in libs.items()) + "."])]

    # ---- 8 CI/CD
    c += [H1("8. Automated testing and CI/CD"),
          P(f"The tests/ folder contains {n_tests} pytest cases covering data cleaning, the "
            "feature transformer, training/evaluation, saved-model reproducibility, the "
            "predictor, MLflow logging and the full HTTP contract of the API (valid and invalid "
            "payloads, metrics endpoint, fail-fast startup). Tests use a synthetic "
            "Cleveland-shaped dataset so CI does not depend on the network."),
          TAB(["Job", "What it does", "Evidence"], [
              ["lint", "ruff static analysis", "annotations on the PR / step log"],
              ["test", "pytest + coverage, JUnit + coverage XML", "test-reports, coverage-report"],
              ["train", "download -> clean -> EDA -> train/track", "model, mlruns, reports"],
              ["docker", "build image, run, smoke-test, helm lint, "
                         "kustomize build", "container-log artifact"],
              ["publish", "push image to GHCR (main only)", "package page"]],
              [0.9, 3.6, 1.9]),
          P("Jobs depend on each other (needs:), so a lint, test, training or container failure "
            "stops the workflow and nothing is published. A deliberately broken run is "
            "included below as proof."),
          SHOT("ci_workflow_overview.png"), SHOT("ci_workflow_artifacts.png"),
          SHOT("ci_failure_demo.png")]

    # ---- 9 container
    c += [H1("9. Model containerisation"),
          P("The Dockerfile starts from python:3.11-slim, installs only requirements-serve.txt "
            "(no MLflow/matplotlib, keeping the image small), copies the source and the model "
            "package, runs as a non-root user (uid 10001) and defines a HEALTHCHECK."),
          TAB(["Endpoint", "Purpose"], [
              ["POST /predict", "one patient -> prediction, confidence, probability, risk level"],
              ["POST /predict/batch", "up to 100 patients"],
              ["GET /health, /ready", "liveness / readiness (ready = model loaded)"],
              ["GET /model-info", "name, version, training metrics"],
              ["GET /metrics", "Prometheus exposition"]], [1.8, 4.6]),
          P("Input is validated by Pydantic (types, clinical ranges, allowed codes, unknown "
            "fields rejected); invalid requests return HTTP 422. Example exchange:"),
          CODE(example),
          SHOT("docker_build.png"), SHOT("docker_run_predict.png"), SHOT("swagger_ui.png")]

    # ---- 10 deployment
    c += [H1("10. Production deployment (Kubernetes)"),
          P("The container is deployed to a local Kubernetes cluster (Docker Desktop or "
            "Minikube). Both plain manifests (k8s/, applied with kustomize) and an equivalent "
            "Helm chart (helm/heart-risk-api) are provided; the same files work on GKE, EKS "
            "or AKS after pushing the image to a registry."),
          TAB(["Resource", "Purpose"], [
              ["Deployment (2 replicas)", "rolling updates (maxUnavailable 0), startup/readiness/"
                                          "liveness probes, resource requests/limits, non-root, "
                                          "read-only root FS, dropped capabilities"],
              ["Service type LoadBalancer", "public endpoint (localhost:80 on Docker Desktop; "
                                            "minikube tunnel on Minikube)"],
              ["Ingress (nginx)", "host-based routing (heart-risk.local)"],
              ["HorizontalPodAutoscaler", "2-5 replicas at 70 % CPU"],
              ["PodDisruptionBudget", "keeps at least one pod during node drains"]],
              [2.0, 4.4]),
          CODE("make k8s-deploy                         # or: make helm-install\n"
               "BASE_URL=http://localhost bash scripts/smoke_test.sh"),
          SHOT("k8s_pods_services.png"), SHOT("k8s_predict_via_lb.png"),
          SHOT("k8s_ingress.png")]

    # ---- 11 monitoring
    c += [H1("11. Monitoring and logging"),
          P("Logging: every request emits one structured JSON line (timestamp, request_id, "
            "method, path, status, latency, client) plus a 'prediction' event with the "
            "outcome. Raw patient features are only logged at DEBUG level to limit sensitive "
            "data exposure. Logs go to stdout (kubectl logs) and optionally to a file."),
          TAB(["Prometheus metric", "Meaning"], [
              ["heart_api_requests_total{method,path,status}", "traffic and error rate"],
              ["heart_api_request_duration_seconds", "latency histogram (p50/p95/p99)"],
              ["heart_api_requests_in_progress", "concurrency"],
              ["heart_api_predictions_total{label}", "prediction mix (concept-drift hint)"],
              ["heart_api_prediction_confidence / _probability", "model confidence distribution"],
              ["heart_api_input_age_years", "input-drift signal vs training mean age"],
              ["heart_api_model_info{version}", "which model version is serving"]],
              [3.2, 3.2]),
          P("Prometheus discovers pods through Kubernetes annotations; Grafana is pre-provisioned "
            "with a 12-panel dashboard (availability, request rate, latency percentiles, 4xx/5xx, "
            "prediction mix, confidence, drift). scripts/generate_traffic.py produces load, "
            "including deliberately invalid requests."),
          SHOT("prometheus_targets.png"), SHOT("grafana_dashboard.png"), SHOT("api_logs.png")]

    # ---- 12 conclusions
    c += [H1("12. Limitations and future work"),
          BUL(["Small single-centre dataset: performance estimates have wide confidence "
               "intervals and will not transfer automatically to other populations.",
               "The 0.5 decision threshold is a default; a clinical deployment would choose a "
               "threshold from the recall/precision trade-off and calibrate probabilities.",
               "Monitoring detects service health and simple drift signals; automatic "
               "retraining triggers and a model registry stage-gate are natural next steps.",
               "Add authentication, rate limiting and TLS at the Ingress before real exposure."]),
          H1("13. Access and deliverables"),
          TAB(["Deliverable", "Location"], [
              ["Code repository", a.repo_url],
              ["Deployed API / access instructions", a.api_url],
              ["Pipeline video", a.video_url],
              ["Notebooks", "notebooks/01_eda, 02_training, 03_inference"],
              ["Tests / CI", "tests/, .github/workflows/ci.yml"],
              ["Manifests / Helm", "k8s/, helm/heart-risk-api/"],
              ["Screenshots", "screenshots/"]], [2.4, 4.0])]
    return c


# ------------------------------------------------------------------ small helpers
def _count_tests() -> int:
    return sum(len(re.findall(r"^def test_", p.read_text(), re.M))
               for p in (ROOT / "tests").glob("test_*.py"))


def _top_correlations() -> str:
    try:
        import pandas as pd

        from heartrisk import config

        df = pd.read_csv(ROOT / "data" / "processed" / "heart_clean.csv")
        corr = df.corr(numeric_only=True)[config.TARGET].drop(config.TARGET)
        top = corr.reindex(corr.abs().sort_values(ascending=False).index).head(5)
        return ", ".join(f"{k} (r = {v:+.2f})" for k, v in top.items())
    except Exception:
        return "see reports/eda_summary.md"


def _api_example() -> str:
    sample = json.loads((ROOT / "scripts" / "sample_request.json").read_text())
    req = "POST /predict\n" + json.dumps(sample)
    try:
        from heartrisk.predictor import HeartRiskPredictor

        pred = HeartRiskPredictor(ROOT / "models" / "model.joblib")
        out = pred.predict([sample])[0]
        out.update({"model_version": pred.version, "request_id": "<uuid>"})
        return req + "\n\n200 OK\n" + json.dumps(out, indent=2)
    except Exception:
        return req + "\n\n(run `make train` to include a live response here)"


# ------------------------------------------------------------------ renderers
def render_markdown(blocks, out: Path, a) -> None:
    md = ["# Heart-Disease Risk Prediction - MLOps Assignment 01 Report\n",
          f"**Course:** MLOps (AIMLCZG523) &nbsp; **Author:** {a.author} ({a.student_id}) "
          f"&nbsp; **Date:** {dt.date.today():%d %B %Y}\n"]
    for b in blocks:
        k = b[0]
        if k == "h1": md.append(f"\n## {b[1]}\n")
        elif k == "h2": md.append(f"\n### {b[1]}\n")
        elif k == "p": md.append(b[1] + "\n")
        elif k == "bul": md.append("\n".join(f"* {i}" for i in b[1]) + "\n")
        elif k == "code": md.append("```\n" + b[1] + "\n```\n")
        elif k == "note": md.append(f"> *{b[1]}*\n")
        elif k == "img":
            rel = Path(b[1]).relative_to(ROOT)
            md.append(f"![{b[2]}](../{rel.as_posix()})\n\n*Figure: {b[2]}*\n")
        elif k == "tab":
            md.append("| " + " | ".join(b[1]) + " |\n|" + "---|" * len(b[1]))
            md.extend("| " + " | ".join(str(x) for x in r) + " |" for r in b[2])
            md.append("")
    out.write_text("\n".join(md))


def render_docx(blocks, out: Path, a) -> None:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Inches, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(1.9)

    base = doc.styles["Normal"]
    base.font.name, base.font.size = "Calibri", Pt(10.5)
    base.paragraph_format.space_after = Pt(5)
    for name, size, color in (("Heading 1", 15, "264653"), ("Heading 2", 12, "2A9D8F")):
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.bold = "Calibri", Pt(size), True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
        st.paragraph_format.keep_with_next = True

    def shade(cell_or_par, hex_fill, is_cell=True):
        el = cell_or_par._tc.get_or_add_tcPr() if is_cell else cell_or_par._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_fill)
        el.append(shd)

    # --- title block
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = t.add_run("Heart-Disease Risk Prediction"); r.bold = True; r.font.size = Pt(24)
    r.font.color.rgb = RGBColor.from_string("264653")
    s = doc.add_paragraph(); r = s.add_run("End-to-End MLOps: model development, CI/CD, deployment "
                                           "and monitoring"); r.font.size = Pt(13)
    r.font.color.rgb = RGBColor.from_string("2A9D8F")
    meta = doc.add_paragraph()
    meta.add_run("Course: ").bold = True; meta.add_run("MLOps (AIMLCZG523) - Assignment 01\n")
    meta.add_run("Author: ").bold = True; meta.add_run(f"{a.author}  |  {a.student_id}\n")
    meta.add_run("Date: ").bold = True; meta.add_run(f"{dt.date.today():%d %B %Y}\n")
    meta.add_run("Repository: ").bold = True; meta.add_run(a.repo_url)

    for b in blocks:
        k = b[0]
        if k == "h1": doc.add_heading(b[1], level=1)
        elif k == "h2": doc.add_heading(b[1], level=2)
        elif k == "p": doc.add_paragraph(b[1])
        elif k == "bul":
            for item in b[1]:
                doc.add_paragraph(item, style="List Bullet")
        elif k == "code":
            for line in b[1].split("\n"):
                p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.left_indent = Cm(0.3)
                run = p.add_run(line or " "); run.font.name = "Consolas"; run.font.size = Pt(8.5)
                shade(p, "F1F3F5", is_cell=False)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif k == "note":
            p = doc.add_paragraph(); run = p.add_run(b[1]); run.italic = True
            run.font.color.rgb = RGBColor.from_string("B5651D"); shade(p, "FFF4E5", is_cell=False)
        elif k == "img":
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(b[1]), width=Inches(min(b[3], 6.6)))
            cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cr = cap.add_run(f"Figure: {b[2]}"); cr.italic = True; cr.font.size = Pt(9)
        elif k == "tab":
            head, rows, widths = b[1], b[2], b[3]
            tbl = doc.add_table(rows=1, cols=len(head)); tbl.style = "Table Grid"
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, h in enumerate(head):
                cell = tbl.rows[0].cells[i]; cell.text = ""
                run = cell.paragraphs[0].add_run(h); run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255); run.font.size = Pt(9.5)
                shade(cell, "264653")
            for ri, row in enumerate(rows):
                cells = tbl.add_row().cells
                for i, val in enumerate(row):
                    cells[i].text = ""
                    run = cells[i].paragraphs[0].add_run(str(val)); run.font.size = Pt(9)
                    if ri % 2: shade(cells[i], "EEF5F4")
            if widths:
                for row in tbl.rows:
                    for i, w in enumerate(widths):
                        row.cells[i].width = Inches(w)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # --- footer page numbers
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run("Heart-Disease MLOps · page ").font.size = Pt(8.5)
    run = fp.add_run(); run.font.size = Pt(8.5)
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
        run._r.append(el)
    doc.save(out)


def to_pdf(docx: Path) -> bool:
    exe = shutil.which("soffice") or shutil.which("libreoffice") or \
        "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    if not exe or not Path(exe).exists():
        print("PDF skipped: LibreOffice not found (brew install --cask libreoffice), "
              "or open the .docx in Word and 'Save as PDF'.")
        return False
    subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", str(docx.parent),
                    str(docx)], check=True, capture_output=True)
    return True


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawTextHelpFormatter)
    ap.add_argument("--author", default="<Your Name>")
    ap.add_argument("--student-id", default="<Student ID>")
    ap.add_argument("--repo-url", default="https://github.com/<your-username>/heart-disease-mlops")
    ap.add_argument("--api-url", default="http://localhost (Docker Desktop LoadBalancer) - see "
                                         "README 'Access'")
    ap.add_argument("--video-url", default="<link to recorded pipeline walkthrough>")
    ap.add_argument("--out", default=str(ROOT / "reports" / "Heart_Disease_MLOps_Report"))
    ap.add_argument("--pdf", action="store_true", help="also export PDF via LibreOffice")
    a = ap.parse_args()

    blocks = build_content(a)
    out = Path(a.out); out.parent.mkdir(parents=True, exist_ok=True)
    render_markdown(blocks, out.with_suffix(".md"), a)
    render_docx(blocks, out.with_suffix(".docx"), a)
    print(f"wrote {out.with_suffix('.md')}\nwrote {out.with_suffix('.docx')}")
    if a.pdf and to_pdf(out.with_suffix(".docx")):
        print(f"wrote {out.with_suffix('.pdf')}")
    if MISSING:
        print(f"\n{len(MISSING)} figure(s)/screenshot(s) still pending:")
        for m in MISSING:
            print("  -", m)
    return 0


if __name__ == "__main__":
    sys.exit(main())
